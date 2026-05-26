from __future__ import annotations

import os
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, jsonify, redirect, request, send_file, send_from_directory
from openpyxl import load_workbook
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BASE_DIR / "frontend"
DB_PATH = BASE_DIR / "temps.db"
REQUIREMENTS_PATH = BASE_DIR / "temperature_requirements.xlsx"

app = Flask(__name__, static_folder=None)


def as_float(value):
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def classify_temperature(actual_temp, requirement, tolerance=2.0):
    actual = as_float(actual_temp)
    required = as_float(requirement)
    if actual is None or required is None:
        return "CRITICAL"
    return "CRITICAL" if actual - required >= tolerance else "OK"


def normalize_lookup_key(value):
    if value is None:
        return None
    text = str(value).strip().upper()
    text = re.sub(r"[^A-Z0-9/:-]", "", text)
    return text or None


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def find_requirement_column(headers):
    aliases = {
        "requirement",
        "required",
        "requiredtemp",
        "requiredtemperature",
        "setpoint",
        "setpointtemp",
        "target",
        "targettemp",
    }
    for index, header in enumerate(headers):
        if normalize_header(header) in aliases:
            return index
    return None


def find_code_columns(headers):
    aliases = {"unitcode", "sourcecode", "baseroom", "roomcode", "room", "roomid"}
    return [index for index, header in enumerate(headers) if normalize_header(header) in aliases]


def read_requirements_workbook(source):
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}, 0

    headers = [str(value or "").strip() for value in rows[0]]
    requirement_index = find_requirement_column(headers)
    code_indexes = find_code_columns(headers)
    if requirement_index is None:
        raise ValueError("Excel file needs a Requirement, Setpoint, or Target column")
    if not code_indexes:
        raise ValueError("Excel file needs a Unit Code, Source Code, Base Room, or Room Code column")

    requirements = {}
    row_count = 0
    for row in rows[1:]:
        if requirement_index >= len(row):
            continue
        requirement = as_float(row[requirement_index])
        if requirement is None:
            continue

        matched = False
        for index in code_indexes:
            key = normalize_lookup_key(row[index] if index < len(row) else None)
            if key:
                requirements[key] = requirement
                matched = True
        if matched:
            row_count += 1

    return requirements, row_count


def load_temperature_requirements():
    if not REQUIREMENTS_PATH.exists():
        return {}

    try:
        requirements, _ = read_requirements_workbook(REQUIREMENTS_PATH)
        return requirements
    except Exception as exc:
        print(f"Temperature requirements workbook error: {exc}")
        return {}


def get_temperature_requirement(room, requirements):
    for field in ("unit_code", "source_code", "base_room"):
        key = normalize_lookup_key(room.get(field))
        if key and key in requirements:
            return requirements[key]
    return as_float(room.get("Requirement"))


def db_last_synced():
    timestamps = []
    for path in (DB_PATH, REQUIREMENTS_PATH):
        if path.exists():
            timestamps.append(datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat())
    return max(timestamps, default=None)


@app.after_request
def set_cache_headers(response):
    path = getattr(response, "_request_path", None)
    if path and path.endswith((".html", ".js", ".css")):
        response.cache_control.no_store = True
        response.cache_control.max_age = 0
    return response


@app.route("/")
def root():
    return redirect("/Temperature/index.html")


@app.route("/Temperature/<path:file_path>")
def temperature_assets(file_path):
    response = send_from_directory(FRONTEND_DIR / "Temperature", file_path)
    response._request_path = file_path
    return response


@app.route("/shared/<path:file_path>")
def shared_assets(file_path):
    response = send_from_directory(FRONTEND_DIR / "shared", file_path)
    response._request_path = file_path
    return response


@app.route("/static/<path:file_path>")
def static_assets(file_path):
    return send_from_directory(FRONTEND_DIR / "static", file_path)


@app.route("/api/page-sync/<page_key>")
def page_sync(page_key):
    return jsonify({
        "page": page_key,
        "last_synced": db_last_synced(),
    })


@app.route("/api/temperature/rooms")
def temperature_rooms():
    if not DB_PATH.exists():
        return jsonify([])

    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute("SELECT * FROM room_temperature").fetchall()
        conn.close()
    except sqlite3.Error as exc:
        print(f"Temperature DB error: {exc}")
        return jsonify([])

    requirements = load_temperature_requirements()
    room_groups = {}
    for row in rows:
        room = dict(row)
        actual = as_float(room.get("Actual Temp"))
        required = get_temperature_requirement(room, requirements)
        diff = None if actual is None or required is None else actual - required
        status = classify_temperature(actual, required)
        max_normal = required + 2 if required is not None else None
        expected_range = {
            "min_normal": None,
            "max_normal": max_normal,
            "attention_buffer": None,
            "warning_buffer": 2,
            "label": f"Required {required:.2f} deg C, red at +2.00 deg C" if required is not None else None,
            "area_group": "Source temperature requirement",
            "configured": actual is not None and required is not None,
        }
        unit = {
            "unit_code": room.get("unit_code") or room.get("base_room"),
            "unit_label": room.get("unit_label") or room.get("room_name") or room.get("base_room"),
            "source_code": room.get("source_code") or room.get("unit_code") or room.get("base_room"),
            "Actual Temp": actual,
            "Requirement": required,
            "temp_diff": diff,
            "status": status,
            "expected_range": expected_range,
        }
        room_key = room.get("base_room") or unit["unit_code"]
        grouped = room_groups.setdefault(room_key, {
            **room,
            "status": status,
            "temperature_status": "NORMAL" if status == "OK" else "WARNING",
            "temperature_threshold_status": "NORMAL" if status == "OK" else "WARNING",
            "expected_range": expected_range,
            "energy": {"mapped": False, "available": False},
            "combined_insight": None,
            "Actual Temp": actual,
            "Requirement": required,
            "temp_diff": diff,
            "units": [],
        })
        grouped["units"].append(unit)

        if status == "CRITICAL" or grouped.get("Actual Temp") is None:
            grouped["status"] = "CRITICAL"
            grouped["temperature_status"] = "WARNING"
            grouped["temperature_threshold_status"] = "WARNING"
            grouped["Actual Temp"] = actual
            grouped["Requirement"] = required
            grouped["temp_diff"] = diff
            grouped["expected_range"] = expected_range

    return jsonify(list(room_groups.values()))


@app.route("/api/temperature/requirements", methods=["GET", "POST"])
def temperature_requirements_file():
    if request.method == "GET":
        if not REQUIREMENTS_PATH.exists():
            return jsonify({"exists": False, "error": "Temperature requirements workbook missing"}), 404

        return send_file(
            REQUIREMENTS_PATH,
            as_attachment=True,
            download_name="temperature_requirements.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Upload an Excel requirements file"}), 400

    filename = secure_filename(uploaded.filename)
    if not filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Requirements file must be .xlsx"}), 400

    try:
        _, row_count = read_requirements_workbook(uploaded)
        uploaded.seek(0)
        uploaded.save(REQUIREMENTS_PATH)
    except Exception as exc:
        return jsonify({"error": f"Could not read requirements workbook: {exc}"}), 400

    return jsonify({
        "ok": True,
        "file": "temperature_requirements.xlsx",
        "rows": row_count,
    })


@app.route("/api/export/report")
def export_report_placeholder():
    return jsonify({"error": "Report export is not included in this standalone temperature clone."}), 501


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5001"))
    app.run(host="127.0.0.1", port=port)
